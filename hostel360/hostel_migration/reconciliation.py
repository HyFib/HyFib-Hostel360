"""Migration reconciliation: does Hostel360 now hold what the legacy system held?

Three checks, computed before a dry run is rolled back so a rehearsal produces the same
report as the real thing:

* **Counts** per entity: rows in the source file (less the rows the mapping excludes)
  against Hostel360 records that came from that entity.
* **Money**: the mapping's ``reconcile`` totals (fee invoices total and outstanding,
  payments, mess bills) summed from the source column and from the target field.
* **Attachments**: SHA-256 of every legacy file against the bytes stored in Hostel360.

The result is stored as JSON on the Migration Run and also written as Markdown (for the
sign-off note) and Excel (for the institution's accounts section).
"""

import io
import json
from decimal import Decimal

import frappe
from frappe import _
from frappe.query_builder.functions import Count, Sum
from frappe.utils import flt, now_datetime

from hostel360.hostel_migration.engine import (
	ATTACHMENT_CHANGED,
	ATTACHMENT_DRY_RUN,
	ATTACHMENT_VERIFIED,
	MAP_DOCTYPE,
)

MONEY_TOLERANCE = Decimal("0.005")


def _join_query(source_system: str, entity_name: str, doctype: str):
	record_map = frappe.qb.DocType(MAP_DOCTYPE)
	target = frappe.qb.DocType(doctype)
	query = (
		frappe.qb.from_(record_map)
		.join(target)
		.on(target.name == record_map.target_name)
		.where(record_map.source_system == source_system)
		.where(record_map.source_entity == entity_name)
		.where(record_map.target_doctype == doctype)
	)
	meta = frappe.get_meta(doctype)
	if meta.is_submittable:
		query = query.where(target.docstatus < 2)
	return query, target


def target_count(source_system: str, entity_name: str, doctype: str) -> int:
	query, _target = _join_query(source_system, entity_name, doctype)
	result = query.select(Count(frappe.qb.DocType(MAP_DOCTYPE).name)).run()
	return int(result[0][0] or 0)


def target_total(source_system: str, entity_name: str, doctype: str, fieldname: str) -> Decimal:
	query, target = _join_query(source_system, entity_name, doctype)
	result = query.select(Sum(target[fieldname])).run()
	return Decimal(str(result[0][0] or 0)).quantize(Decimal("0.01"))


def money(value: Decimal) -> float:
	return float(Decimal(value).quantize(Decimal("0.01")))


def reconcile(engine) -> dict:
	system = engine.mapping.source_system
	entities = []
	totals = []
	all_ok = True

	for stats in engine.stats.values():
		entity = engine.mapping.entities[stats.entity]
		count = 0
		if stats.status != "Failed":
			count = target_count(system, entity.name, stats.target_doctype)
		stats.target_count = count
		expected = stats.source_rows - stats.excluded
		matched = stats.status != "Failed" and stats.failed == 0 and count == expected
		all_ok = all_ok and matched
		entities.append(
			{
				"entity": entity.name,
				"target_doctype": stats.target_doctype,
				"source_file": stats.source_file,
				"source_checksum": stats.source_checksum,
				"source_rows": stats.source_rows,
				"excluded": stats.excluded,
				"expected_records": expected,
				"target_records": count,
				"created": stats.created,
				"updated": stats.updated,
				"unchanged": stats.skipped,
				"failed": stats.failed,
				"doctype_total": frappe.db.count(stats.target_doctype) if stats.status != "Failed" else None,
				"matched": matched,
			}
		)

		for total in entity.totals:
			source_value = stats.source_totals.get(total.label, Decimal("0"))
			target_value = Decimal("0")
			if stats.status != "Failed":
				target_value = target_total(system, entity.name, entity.doctype, total.target_field)
			difference = (target_value - source_value).quantize(Decimal("0.01"))
			unparsed = stats.unparsed_totals.get(total.label, 0)
			ok = abs(difference) < MONEY_TOLERANCE and not unparsed
			all_ok = all_ok and ok
			totals.append(
				{
					"entity": entity.name,
					"label": total.label,
					"source_column": total.source_column,
					"target_field": f"{entity.doctype}.{total.target_field}",
					"source_total": money(source_value),
					"target_total": money(target_value),
					"difference": money(difference),
					"unparsed_source_values": unparsed,
					"matched": ok,
				}
			)

	verified = sum(1 for a in engine.attachments if a["status"] == ATTACHMENT_VERIFIED)
	checked = sum(1 for a in engine.attachments if a["status"] == ATTACHMENT_DRY_RUN)
	changed = sum(1 for a in engine.attachments if a["status"] == ATTACHMENT_CHANGED)
	mismatched = [
		a
		for a in engine.attachments
		if a["status"] not in (ATTACHMENT_VERIFIED, ATTACHMENT_DRY_RUN, ATTACHMENT_CHANGED)
	]
	all_ok = all_ok and not mismatched

	return {
		"run": engine.run_name,
		"mapping": engine.mapping.name,
		"source_system": system,
		"mapping_checksum": engine.mapping.checksum,
		"dry_run": engine.dry_run,
		"prepared_on": now_datetime().strftime("%Y-%m-%d %H:%M:%S"),
		"site": frappe.local.site,
		"status": "Matched" if all_ok else "Mismatched",
		"entities": entities,
		"totals": totals,
		"attachments": {
			"total": len(engine.attachments),
			"verified": verified,
			"checked_dry_run": checked,
			"changed_since_migration": changed,
			"mismatched": len(mismatched),
			"files": engine.attachments,
		},
		"failures": {"count": len(engine.failures), "first": engine.failures[:200]},
	}


# ----------------------------------------------------------------------------- output


def _yes_no(value: bool) -> str:
	return _("Yes") if value else _("No")


def _md_cell(value) -> str:
	text = "" if value is None else str(value)
	return text.replace("|", "\\|").replace("\n", " ")


def _md_table(headers: list[str], rows: list[list]) -> list[str]:
	lines = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _h in headers) + "|"]
	lines.extend("| " + " | ".join(_md_cell(v) for v in row) + " |" for row in rows)
	return lines


def render_markdown(reconciliation: dict, failures: list[dict]) -> str:
	from hostel360.utils.jinja import format_inr

	r = reconciliation
	lines = [
		f"# {_('Migration Reconciliation')}: {r['run']}",
		"",
		f"- {_('Mapping')}: {r['mapping']} (`{r['mapping_checksum'][:12]}`)",
		f"- {_('Source system')}: {r.get('source_system') or ''}",
		f"- {_('Site')}: {r['site']}",
		f"- {_('Prepared on')}: {r['prepared_on']}",
		f"- {_('Dry run')}: {_yes_no(r['dry_run'])}",
		f"- **{_('Result')}: {_(r['status'])}**",
		"",
		f"## {_('Record counts')}",
		"",
	]
	lines += _md_table(
		[_("Entity"), _("DocType"), _("Source rows"), _("Excluded"), _("Expected"), _("In Hostel360"), _("Failed"), _("Matched")],
		[
			[
				e["entity"],
				e["target_doctype"],
				e["source_rows"],
				e["excluded"],
				e["expected_records"],
				e["target_records"],
				e["failed"],
				_yes_no(e["matched"]),
			]
			for e in r["entities"]
		],
	)
	if r["totals"]:
		lines += ["", f"## {_('Money totals')}", ""]
		lines += _md_table(
			[_("Entity"), _("Total"), _("Legacy"), _("Hostel360"), _("Difference"), _("Matched")],
			[
				[
					t["entity"],
					t["label"],
					format_inr(t["source_total"]),
					format_inr(t["target_total"]),
					format_inr(t["difference"]),
					_yes_no(t["matched"]),
				]
				for t in r["totals"]
			],
		)
	attachments = r["attachments"]
	lines += [
		"",
		f"## {_('Attachments')}",
		"",
		f"- {_('Files')}: {attachments['total']}",
		f"- {_('Verified by SHA-256')}: {attachments['verified']}",
		f"- {_('Checked in dry run')}: {attachments['checked_dry_run']}",
		f"- {_('Replaced in Hostel360 since migration')}: {attachments.get('changed_since_migration', 0)}",
		f"- {_('Mismatched')}: {attachments['mismatched']}",
	]
	if failures:
		lines += ["", f"## {_('Rows that failed')} ({len(failures)})", ""]
		lines += _md_table(
			[_("Entity"), _("Line"), _("Legacy ID"), _("Problem")],
			[[f["entity"], f["line"] or "", f["source_id"] or "", f["error"]] for f in failures[:500]],
		)
		if len(failures) > 500:
			lines += ["", _("The Excel report lists all {0} failed rows.").format(len(failures))]
	lines += [
		"",
		f"## {_('Source files')}",
		"",
	]
	lines += _md_table(
		[_("Entity"), _("File"), "SHA-256"],
		[[e["entity"], e["source_file"], e["source_checksum"] or ""] for e in r["entities"]],
	)
	return "\n".join(lines) + "\n"


def build_xlsx(reconciliation: dict, failures: list[dict]) -> bytes:
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill

	r = reconciliation
	workbook = Workbook()
	bold = Font(bold=True)
	header_fill = PatternFill("solid", fgColor="DDE7F3")
	bad_fill = PatternFill("solid", fgColor="F8D7DA")

	def sheet(title: str, headers: list[str], rows: list[list], flag_index: int | None = None):
		ws = workbook.create_sheet(title)
		ws.append(headers)
		for cell in ws[1]:
			cell.font = bold
			cell.fill = header_fill
		for row in rows:
			ws.append(row)
			if flag_index is not None and row[flag_index] == _("No"):
				for cell in ws[ws.max_row]:
					cell.fill = bad_fill
		for idx, header in enumerate(headers, start=1):
			width = max([len(str(header))] + [len(str(row[idx - 1] or "")) for row in rows[:500]])
			ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(width + 2, 10), 60)
		ws.freeze_panes = "A2"
		return ws

	summary = workbook.active
	summary.title = _("Summary")
	for label, value in (
		(_("Migration Run"), r["run"]),
		(_("Mapping"), r["mapping"]),
		(_("Source system"), r.get("source_system")),
		(_("Mapping checksum"), r["mapping_checksum"]),
		(_("Site"), r["site"]),
		(_("Prepared on"), r["prepared_on"]),
		(_("Dry run"), _yes_no(r["dry_run"])),
		(_("Result"), _(r["status"])),
		(_("Failed rows"), len(failures)),
		(_("Attachments verified"), r["attachments"]["verified"]),
	):
		summary.append([label, value])
		summary.cell(row=summary.max_row, column=1).font = bold
	summary.column_dimensions["A"].width = 24
	summary.column_dimensions["B"].width = 70

	sheet(
		_("Counts"),
		[_("Entity"), _("DocType"), _("Source File"), _("Source Rows"), _("Excluded"), _("Expected"), _("In Hostel360"), _("Created"), _("Updated"), _("Unchanged"), _("Failed"), _("Matched")],
		[
			[
				e["entity"],
				e["target_doctype"],
				e["source_file"],
				e["source_rows"],
				e["excluded"],
				e["expected_records"],
				e["target_records"],
				e["created"],
				e["updated"],
				e["unchanged"],
				e["failed"],
				_yes_no(e["matched"]),
			]
			for e in r["entities"]
		],
		flag_index=11,
	)
	sheet(
		_("Money"),
		[_("Entity"), _("Total"), _("Source Column"), _("Target Field"), _("Legacy (INR)"), _("Hostel360 (INR)"), _("Difference"), _("Unparsed Values"), _("Matched")],
		[
			[
				t["entity"],
				t["label"],
				t["source_column"],
				t["target_field"],
				flt(t["source_total"], 2),
				flt(t["target_total"], 2),
				flt(t["difference"], 2),
				t["unparsed_source_values"],
				_yes_no(t["matched"]),
			]
			for t in r["totals"]
		],
		flag_index=8,
	)
	sheet(
		_("Attachments"),
		[_("Entity"), _("Legacy ID"), _("Record"), _("Field"), _("Source File"), _("Source SHA-256"), _("Stored File"), _("Stored SHA-256"), _("Status")],
		[
			[
				a["entity"],
				a["source_id"],
				f"{a['doctype']} {a['name']}",
				a["field"],
				a["source_path"],
				a["source_sha256"],
				a["target_file"],
				a["target_sha256"],
				a["status"],
			]
			for a in r["attachments"]["files"]
		],
	)
	sheet(
		_("Failed Rows"),
		[_("Entity"), _("Line"), _("Legacy ID"), _("Problem")],
		[[f["entity"], f["line"], f["source_id"], f["error"]] for f in failures],
	)

	buffer = io.BytesIO()
	workbook.save(buffer)
	return buffer.getvalue()


def _save_file(run, file_name: str, content: bytes) -> str:
	file_doc = frappe.get_doc(
		{
			"doctype": "File",
			"file_name": file_name,
			"content": content,
			"is_private": 1,
			"attached_to_doctype": run.doctype,
			"attached_to_name": run.name,
		}
	)
	file_doc.flags.ignore_permissions = True
	# Reports are written by the system itself, not uploaded by a person.
	file_doc.flags.hostel360_trusted_file = True
	file_doc.insert()
	return file_doc.file_url


def attach_reports(run, reconciliation: dict, failures: list[dict]) -> None:
	"""Write the JSON, Markdown and Excel reports and link them on the run."""
	stamp = now_datetime().strftime("%Y%m%d-%H%M%S")
	base = f"{run.name}-reconciliation-{stamp}"
	full = {**reconciliation, "failures": {"count": len(failures), "rows": failures}}
	urls = {
		"report_json": _save_file(
			run, f"{base}.json", json.dumps(full, indent=1, default=str, ensure_ascii=False).encode("utf-8")
		),
		"report_markdown": _save_file(run, f"{base}.md", render_markdown(reconciliation, failures).encode("utf-8")),
		"report_file": _save_file(run, f"{base}.xlsx", build_xlsx(reconciliation, failures)),
	}
	run.db_set(urls, update_modified=False)

"""Exit export (tender TR 22): everything the institution needs to leave Hostel360.

The export folder holds

* ``data/csv``   every Hostel360 DocType as CSV, all columns, child tables as their own
  files (with ``parent``/``parentfield``/``idx``), singles as field/value pairs;
* ``data/xlsx``  one workbook per DocType: its records plus a sheet per child table;
* ``data/core``  the Frappe records the data depends on: users and roles, file index,
  audit trail (Version) and comments on Hostel360 records;
* ``database``   a full MariaDB dump of the site (restorable with ``bench restore``);
* ``files``      the site's public and private uploaded files;
* ``manifest.json`` with the size and SHA-256 of every file, and ``README.txt``.

The folder is packed into ``<name>.tar.gz``; with a passphrase the archive is encrypted
(``crypto.py``) and the plaintext copies are removed. ``verify.py`` checks any of these
forms against the manifest.
"""

import csv
import datetime
import decimal
import gzip
import hashlib
import json
import os
import re
import shutil
import subprocess
import tarfile
import tempfile
from shutil import which

import frappe
from frappe import _
from frappe.utils import cstr, now_datetime

from hostel360.hostel_migration import crypto

APP = "hostel360"
FORMAT = "hostel360-exit-export"
FORMAT_VERSION = 1
PAGE_SIZE = 2000
XLSX_CELL_LIMIT = 32767
SHEET_NAME_INVALID = re.compile(r"[\[\]:*?/\\]")
FILE_NAME_INVALID = re.compile(r"[^A-Za-z0-9 ._()-]")
COPY_CHUNK = 1024 * 1024

# Frappe columns that come first in every CSV, in this order.
LEADING_COLUMNS = ("name", "owner", "creation", "modified", "modified_by", "docstatus", "idx")
CHILD_COLUMNS = ("parent", "parentfield", "parenttype")

# Fields of User that are useful to the institution; password hashes live elsewhere and
# API secrets are never exported as data (they are, encrypted, in the database dump).
USER_FIELDS = (
	"name",
	"email",
	"full_name",
	"first_name",
	"last_name",
	"enabled",
	"user_type",
	"mobile_no",
	"language",
	"time_zone",
	"last_login",
	"creation",
	"modified",
)


class ExportError(frappe.ValidationError):
	pass


# ----------------------------------------------------------------------- helpers


def sha256_of(path: str) -> str:
	digest = hashlib.sha256()
	with open(path, "rb") as f:
		while chunk := f.read(COPY_CHUNK):
			digest.update(chunk)
	return digest.hexdigest()


def safe_file_name(name: str) -> str:
	return FILE_NAME_INVALID.sub("_", name).strip() or "unnamed"


def csv_value(value) -> str:
	if value is None:
		return ""
	if isinstance(value, datetime.datetime):
		return value.isoformat(sep=" ")
	if isinstance(value, datetime.date | datetime.time | datetime.timedelta):
		return str(value)
	if isinstance(value, decimal.Decimal):
		return format(value, "f")
	if isinstance(value, bytes):
		return value.hex()
	return cstr(value)


def is_inside(path: str, root: str) -> bool:
	path, root = os.path.realpath(path), os.path.realpath(root)
	return path == root or path.startswith(root + os.sep)


def app_doctypes() -> list[frappe._dict]:
	"""Every DocType of every Hostel360 module, discovered from the site."""
	modules = set(frappe.get_module_list(APP))
	modules.update(frappe.get_all("Module Def", filters={"app_name": APP}, pluck="name"))
	return frappe.get_all(
		"DocType",
		filters={"module": ["in", sorted(modules)], "is_virtual": 0},
		fields=["name", "module", "istable", "issingle", "is_submittable"],
		order_by="module asc, name asc",
		limit_page_length=0,
	)


def ordered_columns(doctype: str, child: bool) -> list[str]:
	db_columns = list(frappe.db.get_table_columns(doctype))
	meta_order = [df.fieldname for df in frappe.get_meta(doctype).fields if df.fieldname in db_columns]
	leading = [c for c in (*LEADING_COLUMNS, *(CHILD_COLUMNS if child else ())) if c in db_columns]
	rest = [c for c in meta_order if c not in leading]
	extra = [c for c in db_columns if c not in leading and c not in rest]
	return leading + rest + extra


def iter_rows(doctype: str, columns: list[str], filters: dict | None = None, order_by=("creation", "name")):
	table = frappe.qb.DocType(doctype)
	offset = 0
	while True:
		query = frappe.qb.from_(table).select(*[table[c] for c in columns])
		for fieldname, value in (filters or {}).items():
			query = query.where(table[fieldname] == value)
		for fieldname in order_by:
			query = query.orderby(table[fieldname])
		rows = query.limit(PAGE_SIZE).offset(offset).run()
		if not rows:
			return
		yield from rows
		if len(rows) < PAGE_SIZE:
			return
		offset += PAGE_SIZE


# ------------------------------------------------------------------------- writers


class XlsxBook:
	"""A write-only workbook that keeps text as text (no formulas) and Excel-safe."""

	def __init__(self):
		from openpyxl import Workbook

		self.book = Workbook(write_only=True)
		self.names: set[str] = set()

	def sheet_name(self, title: str) -> str:
		base = SHEET_NAME_INVALID.sub(" ", title)[:31].strip() or "Sheet"
		name, n = base, 2
		while name.lower() in self.names:
			suffix = f" ({n})"
			name, n = base[: 31 - len(suffix)] + suffix, n + 1
		self.names.add(name.lower())
		return name

	def add_sheet(self, title: str, headers: list[str], rows) -> int:
		from openpyxl.cell import WriteOnlyCell
		from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
		from openpyxl.styles import Font

		sheet = self.book.create_sheet(self.sheet_name(title))
		bold = Font(bold=True)
		header_cells = []
		for header in headers:
			cell = WriteOnlyCell(sheet, value=header)
			cell.font = bold
			header_cells.append(cell)
		sheet.append(header_cells)

		count = 0
		for row in rows:
			cells = []
			for value in row:
				if isinstance(value, decimal.Decimal):
					value = float(value)
				elif isinstance(value, datetime.timedelta):
					value = str(value)
				elif isinstance(value, bytes):
					value = value.hex()
				if isinstance(value, str):
					value = ILLEGAL_CHARACTERS_RE.sub("", value)
					if len(value) > XLSX_CELL_LIMIT:
						value = value[: XLSX_CELL_LIMIT - 20] + " [see CSV]"
					cell = WriteOnlyCell(sheet, value=value)
					# A text value such as "=SUM(...)" stays text, never a live formula.
					cell.data_type = "s"
				else:
					cell = WriteOnlyCell(sheet, value=value)
				cells.append(cell)
			sheet.append(cells)
			count += 1
			if count >= 1_048_000:
				break
		return count

	def save(self, path: str) -> None:
		self.book.save(path)


def write_csv(path: str, headers: list[str], rows) -> int:
	count = 0
	with open(path, "w", newline="", encoding="utf-8-sig") as f:
		writer = csv.writer(f)
		writer.writerow(headers)
		for row in rows:
			writer.writerow([csv_value(v) for v in row])
			count += 1
	return count


# --------------------------------------------------------------------------- export


class ExitExport:
	def __init__(self, output_dir: str, include_files: bool = True, progress=None):
		self.output_dir = os.path.realpath(output_dir)
		self.include_files = include_files
		self.progress = progress
		stamp = now_datetime().strftime("%Y%m%d-%H%M%S")
		site_slug = re.sub(r"[^A-Za-z0-9]+", "-", frappe.local.site).strip("-")
		self.name = f"hostel360-export-{site_slug}-{stamp}"
		self.root = os.path.join(self.output_dir, self.name)
		self.doctypes_summary: list[dict] = []
		self.core_summary: list[dict] = []
		self.database: dict = {}
		self.files: dict = {"included": include_files, "public": 0, "private": 0, "bytes": 0, "skipped": []}
		self.hashes: dict[str, str] = {}

	def say(self, message: str) -> None:
		if self.progress:
			self.progress(message)

	def check_output_dir(self) -> None:
		public = frappe.get_site_path("public")
		if is_inside(self.output_dir, public):
			raise ExportError(_("The export may not be written inside the site's public folder."))
		os.makedirs(self.output_dir, exist_ok=True)
		if os.path.exists(self.root):
			raise ExportError(_("{0} already exists.").format(self.root))
		os.makedirs(self.root, mode=0o700)

	def path(self, *parts: str) -> str:
		full = os.path.join(self.root, *parts)
		os.makedirs(os.path.dirname(full), exist_ok=True)
		return full

	# ------------------------------------------------------------------ data

	def export_data(self) -> None:
		doctypes = app_doctypes()
		children_by_parent: dict[str, list[tuple[str, str, str]]] = {}
		for dt in doctypes:
			if dt.istable or dt.issingle:
				continue
			for df in frappe.get_meta(dt.name).get_table_fields():
				children_by_parent.setdefault(dt.name, []).append((df.fieldname, df.label or df.fieldname, df.options))

		for dt in doctypes:
			self.say(_("Exporting {0}").format(dt.name))
			file_stem = safe_file_name(dt.name)
			entry = {"doctype": dt.name, "module": dt.module, "kind": "document"}
			if dt.issingle:
				entry["kind"] = "single"
				rows = frappe.db.sql(
					"select field, value from `tabSingles` where doctype = %(doctype)s order by field",
					{"doctype": dt.name},
				)
				entry["rows"] = write_csv(self.path("data", "csv", f"{file_stem}.csv"), ["field", "value"], rows)
				book = XlsxBook()
				book.add_sheet(dt.name, ["field", "value"], rows)
				book.save(self.path("data", "xlsx", f"{file_stem}.xlsx"))
				entry.update(csv=f"data/csv/{file_stem}.csv", xlsx=f"data/xlsx/{file_stem}.xlsx")
				self.doctypes_summary.append(entry)
				continue

			child = bool(dt.istable)
			columns = ordered_columns(dt.name, child)
			order_by = ("parenttype", "parent", "parentfield", "idx") if child else ("creation", "name")
			entry["kind"] = "child" if child else "document"
			entry["rows"] = write_csv(
				self.path("data", "csv", f"{file_stem}.csv"), columns, iter_rows(dt.name, columns, order_by=order_by)
			)
			entry["csv"] = f"data/csv/{file_stem}.csv"
			if not child:
				book = XlsxBook()
				book.add_sheet(dt.name, columns, iter_rows(dt.name, columns, order_by=order_by))
				for fieldname, label, child_doctype in children_by_parent.get(dt.name, []):
					child_columns = ordered_columns(child_doctype, True)
					book.add_sheet(
						label,
						child_columns,
						iter_rows(
							child_doctype,
							child_columns,
							filters={"parenttype": dt.name, "parentfield": fieldname},
							order_by=("parent", "idx"),
						),
					)
				book.save(self.path("data", "xlsx", f"{file_stem}.xlsx"))
				entry["xlsx"] = f"data/xlsx/{file_stem}.xlsx"
			self.doctypes_summary.append(entry)

		self.export_core([dt.name for dt in doctypes])

	def export_core(self, doctypes: list[str]) -> None:
		"""Frappe records that the institution's data depends on."""
		user_columns = [c for c in USER_FIELDS if c in frappe.db.get_table_columns("User")]
		users = (row for row in iter_rows("User", user_columns) if row[0] != "Guest")
		self.write_core("User", user_columns, users)

		role_columns = ["parent", "role"]
		self.write_core(
			"Has Role",
			role_columns,
			iter_rows("Has Role", role_columns, filters={"parenttype": "User"}, order_by=("parent", "role")),
		)

		file_columns = [
			c
			for c in (
				"name",
				"file_name",
				"file_url",
				"is_private",
				"file_size",
				"content_hash",
				"attached_to_doctype",
				"attached_to_name",
				"attached_to_field",
				"folder",
				"is_folder",
				"owner",
				"creation",
			)
			if c in frappe.db.get_table_columns("File")
		]
		self.write_core("File", file_columns, iter_rows("File", file_columns))

		for core_doctype, link_field in (("Version", "ref_doctype"), ("Comment", "reference_doctype")):
			columns = ordered_columns(core_doctype, False)
			rows = (
				row
				for doctype in doctypes
				for row in iter_rows(core_doctype, columns, filters={link_field: doctype})
			)
			self.write_core(core_doctype, columns, rows)

	def write_core(self, doctype: str, columns: list[str], rows) -> None:
		stem = safe_file_name(doctype)
		count = write_csv(self.path("data", "core", f"{stem}.csv"), columns, rows)
		self.core_summary.append({"doctype": doctype, "rows": count, "csv": f"data/core/{stem}.csv"})

	# --------------------------------------------------------------- database

	def export_database(self) -> None:
		db_name = frappe.conf.db_name
		target = self.path("database", f"{safe_file_name(db_name)}.sql.gz")
		self.say(_("Dumping the database"))
		binary = which("mariadb-dump")
		if binary and frappe.conf.db_type in (None, "mariadb"):
			self.dump_with_mariadb(binary, target)
			method = "mariadb-dump"
		else:
			self.dump_with_frappe(target)
			method = "frappe.utils.backups"
		self.database = {
			"file": os.path.relpath(target, self.root).replace(os.sep, "/"),
			"method": method,
			"db_name": db_name,
			"db_type": frappe.conf.db_type or "mariadb",
			"restore": "bench --site <site> restore <this file>",
		}

	def dump_with_mariadb(self, binary: str, target: str) -> None:
		conf = frappe.conf
		lines = ["[client]", f"user={_option_value(conf.db_user or conf.db_name)}"]
		if conf.db_password:
			lines.append(f"password={_option_value(conf.db_password)}")
		if conf.db_socket:
			lines.append(f"socket={_option_value(conf.db_socket)}")
		else:
			lines.append(f"host={_option_value(conf.db_host or '127.0.0.1')}")
			lines.append(f"port={int(conf.db_port or 3306)}")

		# Credentials go through a private option file, never the command line where
		# any local user could read them from the process list.
		fd, option_file = tempfile.mkstemp(prefix="h360-dump-", suffix=".cnf")
		error_log = tempfile.TemporaryFile()
		try:
			with os.fdopen(fd, "w") as f:
				f.write("\n".join(lines) + "\n")
			os.chmod(option_file, 0o600)
			command = [
				binary,
				f"--defaults-extra-file={option_file}",
				"--single-transaction",
				"--quick",
				"--lock-tables=false",
				"--hex-blob",
				"--default-character-set=utf8mb4",
				"--triggers",
				conf.db_name,
			]
			with subprocess.Popen(command, stdout=subprocess.PIPE, stderr=error_log) as proc:
				with gzip.open(target, "wb", compresslevel=6) as out:
					out.write(f"-- {FORMAT} database dump of {frappe.local.site}\n".encode())
					while chunk := proc.stdout.read(COPY_CHUNK):
						out.write(chunk)
				returncode = proc.wait()
			if returncode != 0:
				error_log.seek(0)
				detail = error_log.read().decode("utf-8", "replace").strip()
				raise ExportError(_("The database dump failed: {0}").format(detail[-500:] or returncode))
		finally:
			error_log.close()
			os.remove(option_file)

	def dump_with_frappe(self, target: str) -> None:
		from frappe.utils.backups import BackupGenerator

		generator = BackupGenerator(
			frappe.conf.db_name,
			frappe.conf.db_user,
			frappe.conf.db_password,
			db_socket=frappe.conf.db_socket,
			db_host=frappe.conf.db_host,
			db_port=frappe.conf.db_port,
			db_type=frappe.conf.db_type,
			backup_path_db=target,
			ignore_conf=True,
		)
		generator.backup_path_db = target
		generator.take_dump()

	# ----------------------------------------------------------------- files

	def export_files(self) -> None:
		if not self.include_files:
			return
		self.say(_("Copying uploaded files"))
		for kind, source in (("public", frappe.get_site_path("public", "files")), ("private", frappe.get_site_path("private", "files"))):
			source = os.path.realpath(source)
			if not os.path.isdir(source):
				continue
			for folder, dirnames, filenames in os.walk(source, followlinks=False):
				# Never copy the export into itself when it is written below private/files.
				dirnames[:] = [d for d in dirnames if not is_inside(os.path.join(folder, d), self.output_dir)]
				for filename in filenames:
					src = os.path.join(folder, filename)
					if os.path.islink(src) or not os.path.isfile(src):
						self.files["skipped"].append(os.path.relpath(src, source))
						continue
					relative = os.path.relpath(src, source)
					dest = self.path("files", kind, relative)
					digest = hashlib.sha256()
					with open(src, "rb") as fin, open(dest, "wb") as fout:
						while chunk := fin.read(COPY_CHUNK):
							digest.update(chunk)
							fout.write(chunk)
					shutil.copystat(src, dest)
					self.hashes[os.path.relpath(dest, self.root).replace(os.sep, "/")] = digest.hexdigest()
					self.files[kind] += 1
					self.files["bytes"] += os.path.getsize(dest)

	# -------------------------------------------------------------- manifest

	def write_readme(self) -> None:
		institution = frappe.db.get_single_value("Hostel Settings", "institution_name") or ""
		text = f"""HyFib Hostel360 - exit export
Institution : {institution}
Site        : {frappe.local.site}
Created     : {now_datetime().isoformat(sep=" ", timespec="seconds")} (Asia/Kolkata)

Contents
  data/csv/        every Hostel360 table as CSV (UTF-8), one file per DocType
  data/xlsx/       the same data as Excel workbooks, child tables as extra sheets
  data/core/       users, roles, file index, audit trail and comments
  database/        complete MariaDB dump (gzip); restore with: bench --site <site> restore <file>
  files/public/    public uploaded files
  files/private/   private uploaded files (photos, documents, reports)
  manifest.json    size and SHA-256 checksum of every file in this export

Checking the export
  Every file is listed in manifest.json with its SHA-256 checksum. With Hostel360
  installed: bench h360-verify-export --path <this folder or archive>
  Without it: compare "sha256sum <file>" with the manifest entry.

Support: HyFib Solutions Private Limited, support@hyfib.com, +91 96060 60596
"""
		with open(self.path("README.txt"), "w", encoding="utf-8") as f:
			f.write(text)

	def write_manifest(self) -> dict:
		entries = []
		total_bytes = 0
		for folder, _dirs, filenames in os.walk(self.root):
			for filename in sorted(filenames):
				full = os.path.join(folder, filename)
				relative = os.path.relpath(full, self.root).replace(os.sep, "/")
				if relative == "manifest.json":
					continue
				size = os.path.getsize(full)
				digest = self.hashes.get(relative) or sha256_of(full)
				entries.append({"path": relative, "size": size, "sha256": digest})
				total_bytes += size
		entries.sort(key=lambda e: e["path"])

		manifest = {
			"format": FORMAT,
			"format_version": FORMAT_VERSION,
			"export_name": self.name,
			"site": frappe.local.site,
			"created_on": now_datetime().isoformat(timespec="seconds"),
			"created_by": frappe.session.user,
			"time_zone": frappe.utils.get_system_timezone(),
			"institution": frappe.db.get_single_value("Hostel Settings", "institution_name"),
			"apps": {app: _app_version(app) for app in frappe.get_installed_apps()},
			"doctypes": self.doctypes_summary,
			"core": self.core_summary,
			"database": self.database,
			"files": {k: v for k, v in self.files.items() if k != "skipped"},
			"skipped_files": self.files["skipped"],
			"totals": {"entries": len(entries), "bytes": total_bytes},
			"entries": entries,
		}
		with open(self.path("manifest.json"), "w", encoding="utf-8") as f:
			json.dump(manifest, f, indent=1, ensure_ascii=False, default=str)
		return manifest

	# --------------------------------------------------------------- archive

	def make_archive(self) -> str:
		archive = os.path.join(self.output_dir, f"{self.name}.tar.gz")
		self.say(_("Packing {0}").format(os.path.basename(archive)))
		with tarfile.open(archive, "w:gz", compresslevel=6) as tar:
			tar.add(self.root, arcname=self.name, recursive=True)
		os.chmod(archive, 0o600)
		return archive

	def run(self, encrypt: bool = False, passphrase: str | None = None, archive: bool = True) -> dict:
		if encrypt:
			crypto.check_passphrase(passphrase or "")
		self.check_output_dir()
		self.export_data()
		self.export_database()
		self.export_files()
		self.write_readme()
		manifest = self.write_manifest()

		result = {
			"name": self.name,
			"directory": self.root,
			"archive": None,
			"encrypted": False,
			"entries": manifest["totals"]["entries"],
			"bytes": manifest["totals"]["bytes"],
			"doctypes": len(self.doctypes_summary),
			"manifest_sha256": sha256_of(os.path.join(self.root, "manifest.json")),
		}
		if not (archive or encrypt):
			return result

		packed = self.make_archive()
		result["archive"] = packed
		if encrypt:
			self.say(_("Encrypting the archive"))
			encrypted = packed + crypto.ENCRYPTED_SUFFIX
			crypto.encrypt_file(packed, encrypted, passphrase)
			os.remove(packed)
			shutil.rmtree(self.root)
			result.update(archive=encrypted, directory=None, encrypted=True)
		with open(result["archive"] + ".sha256", "w") as f:
			f.write(f"{sha256_of(result['archive'])}  {os.path.basename(result['archive'])}\n")
		result["archive_sha256"] = sha256_of(result["archive"])
		return result


def _option_value(value) -> str:
	"""Quote a value for a MySQL option file."""
	text = cstr(value).replace("\\", "\\\\").replace('"', '\\"')
	return f'"{text}"'


def _app_version(app: str) -> str | None:
	try:
		return frappe.get_attr(f"{app}.__version__")
	except Exception:
		return None


def run_exit_export(
	output_dir: str,
	include_files: bool = True,
	encrypt: bool = False,
	passphrase: str | None = None,
	archive: bool = True,
	progress=None,
) -> dict:
	export = ExitExport(output_dir, include_files=include_files, progress=progress)
	try:
		return export.run(encrypt=encrypt, passphrase=passphrase, archive=archive)
	except Exception:
		# Never leave a half-written export (possibly holding personal data) behind.
		if os.path.isdir(export.root):
			shutil.rmtree(export.root, ignore_errors=True)
		archive_path = os.path.join(export.output_dir, f"{export.name}.tar.gz")
		for leftover in (archive_path, archive_path + crypto.ENCRYPTED_SUFFIX):
			for path in (leftover, leftover + ".sha256"):
				if os.path.exists(path):
					os.remove(path)
		raise

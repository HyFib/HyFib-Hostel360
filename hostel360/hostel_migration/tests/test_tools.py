"""Exit export, demo data, site setup and report rendering."""

import io
import json
import os
import shutil
import tempfile

import frappe
from frappe.tests import IntegrationTestCase, UnitTestCase

from hostel360.hostel_migration.export import ExportError, run_exit_export
from hostel360.hostel_migration.reconciliation import build_xlsx, render_markdown
from hostel360.hostel_migration.site_setup import (
	SiteSetupError,
	clean_portal_url,
	ensure_admin_user,
	generate_password,
)
from hostel360.hostel_migration.verify import verify_export

PASSPHRASE = "export test passphrase 2026"


def sample_reconciliation() -> dict:
	return {
		"run": "MIG-TEST-00001",
		"mapping": "Legacy register",
		"source_system": "legacy",
		"mapping_checksum": "a" * 64,
		"dry_run": True,
		"prepared_on": "2026-10-15 10:00:00",
		"site": "test.localhost",
		"status": "Mismatched",
		"entities": [
			{
				"entity": "students",
				"target_doctype": "Hostel Student",
				"source_file": "students.csv",
				"source_checksum": "b" * 64,
				"source_rows": 10,
				"excluded": 1,
				"expected_records": 9,
				"target_records": 8,
				"created": 8,
				"updated": 0,
				"unchanged": 0,
				"failed": 1,
				"doctype_total": 8,
				"matched": False,
			}
		],
		"totals": [
			{
				"entity": "fee_invoices",
				"label": "Fee invoices total",
				"source_column": "total",
				"target_field": "Hostel Fee Invoice.total_amount",
				"source_total": 125000.0,
				"target_total": 112500.0,
				"difference": -12500.0,
				"unparsed_source_values": 0,
				"matched": False,
			}
		],
		"attachments": {"total": 0, "verified": 0, "checked_dry_run": 0, "mismatched": 0, "files": []},
		"failures": {"count": 1, "first": []},
	}


class TestReportRendering(UnitTestCase):
	failures = [
		{"entity": "students", "line": 7, "source_id": "S6", "error": "mobile_number: '123' | not valid"}
	]

	def test_markdown(self):
		text = render_markdown(sample_reconciliation(), self.failures)
		self.assertIn("# Migration Reconciliation: MIG-TEST-00001", text)
		self.assertIn("**Result: Mismatched**", text)
		self.assertIn("| students | Hostel Student | 10 | 1 | 9 | 8 | 1 | No |", text)
		self.assertIn("₹1,25,000.00", text)
		self.assertIn("-₹12,500.00", text)
		# Pipes inside messages must not break the table.
		self.assertIn("'123' \\| not valid", text)

	def test_excel(self):
		from openpyxl import load_workbook

		workbook = load_workbook(io.BytesIO(build_xlsx(sample_reconciliation(), self.failures)))
		self.assertEqual(workbook.sheetnames, ["Summary", "Counts", "Money", "Attachments", "Failed Rows"])
		money = list(workbook["Money"].iter_rows(values_only=True))
		self.assertEqual(money[1][4:7], (125000.0, 112500.0, -12500.0))
		failed = list(workbook["Failed Rows"].iter_rows(values_only=True))
		self.assertEqual(failed[1][:3], ("students", 7, "S6"))


class TestExitExport(IntegrationTestCase):
	def setUp(self):
		super().setUp()
		self.output = tempfile.mkdtemp(prefix="h360-export-test-")
		self.addCleanup(shutil.rmtree, self.output, True)

	def test_encrypted_export_verifies(self):
		facility = frappe.get_doc(
			{
				"doctype": "Room Facility",
				"facility_name": f"Export Test {frappe.generate_hash(length=6)}",
				"description": '=HYPERLINK("http://example.com","click")',
			}
		).insert(ignore_permissions=True)

		result = run_exit_export(self.output, include_files=False, encrypt=True, passphrase=PASSPHRASE)
		self.assertTrue(result["encrypted"])
		self.assertIsNone(result["directory"])
		self.assertTrue(result["archive"].endswith(".tar.gz.h360enc"))
		self.assertEqual(
			sorted(os.listdir(self.output)),
			sorted([os.path.basename(result["archive"]), os.path.basename(result["archive"]) + ".sha256"]),
		)

		extracted = os.path.join(self.output, "check")
		verified = verify_export(result["archive"], passphrase=PASSPHRASE, extract_to=extracted)
		self.assertTrue(verified.ok, verified.problems)
		root = verified.extracted_to

		with open(os.path.join(root, "manifest.json")) as f:
			manifest = json.load(f)
		exported = {d["doctype"]: d for d in manifest["doctypes"]}
		for doctype in ("Migration Run", "Hostel Student", "Hostel Settings", "Hostel Student Guardian"):
			self.assertIn(doctype, exported)
		self.assertEqual(exported["Hostel Settings"]["kind"], "single")
		self.assertEqual(exported["Hostel Student Guardian"]["kind"], "child")
		self.assertEqual(exported["Room Facility"]["rows"], frappe.db.count("Room Facility"))
		self.assertTrue(manifest["database"]["file"].endswith(".sql.gz"))
		self.assertFalse(manifest["files"]["included"])
		self.assertTrue(all(len(e["sha256"]) == 64 for e in manifest["entries"]))

		# Text that looks like a formula stays text in Excel.
		from openpyxl import load_workbook

		workbook = load_workbook(os.path.join(root, "data", "xlsx", "Room Facility.xlsx"))
		sheet = workbook.worksheets[0]
		headers = [c.value for c in sheet[1]]
		cells = [row[headers.index("description")] for row in sheet.iter_rows(min_row=2)]
		cell = next(c for c in cells if c.value and str(c.value).startswith("=HYPERLINK"))
		self.assertEqual(cell.data_type, "s")
		self.assertIn(facility.name, [row[headers.index("name")].value for row in sheet.iter_rows(min_row=2)])

	def test_plain_export_and_tamper_detection(self):
		result = run_exit_export(self.output, include_files=False, encrypt=False, archive=True)
		self.assertTrue(os.path.isdir(result["directory"]))
		self.assertTrue(verify_export(result["directory"]).ok)
		self.assertTrue(verify_export(result["archive"]).ok)
		with open(os.path.join(result["directory"], "data", "csv", "Hostel.csv"), "a") as f:
			f.write("tampered\n")
		self.assertFalse(verify_export(result["directory"]).ok)

	def test_refuses_weak_passphrase_and_public_folder(self):
		with self.assertRaises(ValueError):
			run_exit_export(self.output, include_files=False, encrypt=True, passphrase="short")
		with self.assertRaises(ExportError):
			run_exit_export(frappe.get_site_path("public", "files", "exports"), include_files=False)
		self.assertEqual(os.listdir(self.output), [])


class TestDemoData(IntegrationTestCase):
	def test_demo_data_is_deterministic_and_idempotent(self):
		from hostel360.hostel_migration.demo_data import generate

		generate(students=6, seed=11)
		students = frappe.get_all(
			"Hostel Student",
			filters={"register_number": ["like", "UAT%"]},
			fields=["name", "email", "status"],
		)
		self.assertGreaterEqual(len(students), 6)
		self.assertTrue(all(s.email.endswith(".invalid") for s in students if s.email))
		self.assertTrue(frappe.db.exists("Hostel", "BH1") and frappe.db.exists("Hostel", "GH1"))
		self.assertTrue(frappe.db.exists("Mess Menu", {"mess": "Boys Mess", "status": "Published"}))

		again = generate(students=6, seed=11)
		self.assertEqual(again["created"], {})
		self.assertEqual(again["errors"], [])

	def test_refuses_while_messaging_is_live(self):
		from hostel360.hostel_migration.demo_data import DemoDataError, generate

		previous = {
			field: frappe.db.get_single_value("Messaging Settings", field)
			for field in ("sms_enabled", "test_mode")
		}
		frappe.db.set_single_value("Messaging Settings", {"sms_enabled": 1, "test_mode": 0})
		frappe.clear_document_cache("Messaging Settings", "Messaging Settings")
		try:
			with self.assertRaises(DemoDataError):
				generate(students=1, seed=1)
		finally:
			frappe.db.set_single_value("Messaging Settings", previous)
			frappe.clear_document_cache("Messaging Settings", "Messaging Settings")


class TestSiteSetup(IntegrationTestCase):
	def test_portal_url(self):
		self.assertEqual(clean_portal_url(" https://hostel.example.ac.in/ "), "https://hostel.example.ac.in")
		self.assertEqual(clean_portal_url("http://b8.localhost:8000"), "http://b8.localhost:8000")
		for bad in (
			"hostel.example.ac.in",
			"http://hostel.example.ac.in",
			"https://x.in/hostel",
			"ftp://x.in",
		):
			with self.assertRaises(SiteSetupError, msg=bad):
				clean_portal_url(bad)

	def test_new_passwords_pass_the_policy(self):
		from frappe.utils.password_strength import test_password_strength

		passwords = {generate_password() for _i in range(20)}
		self.assertEqual(len(passwords), 20)
		for password in passwords:
			self.assertGreaterEqual(test_password_strength(password).get("score", 0), 3)

	def test_admin_user_created_once(self):
		email = f"h360-admin-{frappe.generate_hash(length=8)}@example.com"
		user, password = ensure_admin_user(
			email, "Hostel", "Admin", system_manager=False, reset_password=False
		)
		self.assertTrue(password)
		self.assertEqual(frappe.db.get_value("User", user, "user_type"), "System User")
		self.assertIn("Hostel Administrator", frappe.get_roles(user))
		self.assertNotIn("System Manager", frappe.get_roles(user))

		_user, again = ensure_admin_user(email, "Hostel", "Admin", system_manager=True, reset_password=False)
		self.assertIsNone(again)
		self.assertIn("System Manager", frappe.get_roles(user))

		_user, reset = ensure_admin_user(email, "Hostel", "Admin", system_manager=True, reset_password=True)
		self.assertTrue(reset)
		self.assertNotEqual(reset, password)


class TestCommandHelpers(UnitTestCase):
	def test_paths_relative_to_where_the_command_was_typed(self):
		from hostel360.commands import resolve_cli_path

		folder = tempfile.mkdtemp(prefix="h360-cli-test-")
		self.addCleanup(shutil.rmtree, folder, True)
		os.makedirs(os.path.join(folder, "legacy"))
		previous = os.environ.get("PWD")
		os.environ["PWD"] = folder
		try:
			self.assertEqual(resolve_cli_path("legacy"), os.path.realpath(os.path.join(folder, "legacy")))
			self.assertEqual(resolve_cli_path("/abs/path"), "/abs/path")
		finally:
			if previous is None:
				os.environ.pop("PWD", None)
			else:
				os.environ["PWD"] = previous

	def test_commands_are_registered(self):
		from hostel360.commands import commands

		self.assertEqual(
			sorted(c.name for c in commands),
			[
				"h360-demo-data",
				"h360-demo-reset",
				"h360-demo-site",
				"h360-export",
				"h360-migrate",
				"h360-purge-institution-data",
				"h360-setup",
				"h360-verify-deletion-certificate",
				"h360-verify-export",
			],
		)

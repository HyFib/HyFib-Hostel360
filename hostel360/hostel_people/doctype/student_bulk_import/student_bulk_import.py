# Copyright (c) 2026, HyFib Solutions Private Limited and contributors
# For license information, please see license.txt

"""Bulk registration of students (and their guardians and rooms) from a spreadsheet.

Flow: attach file → Validate (every row is checked, nothing is written) → Start Import
(background job on the "long" queue). Each row is written inside its own savepoint, so
a bad row never leaves half a student behind, and the import can be re-run safely:
students are matched by register number and guardians by mobile number.
"""

import json
import re
from datetime import date, datetime
from io import BytesIO

import frappe
from frappe import _
from frappe.model.document import Document
from frappe.permissions import get_user_permissions
from frappe.utils import getdate, now_datetime, nowdate, validate_email_address
from frappe.utils.background_jobs import is_job_enqueued

from hostel360.hostel_people.doctype.hostel_student.hostel_student import clean_register_number
from hostel360.hostel_people.services import block_restrictions, get_staff_scope
from hostel360.utils.validation import normalize_mobile, validate_pincode

MAX_ROWS = 5000
PREVIEW_ROWS = 200
COMMIT_EVERY = 50

# fieldname, label, required, help text shown on the Instructions sheet
COLUMNS = (
	("register_number", "Register Number", True, "Unique register number, e.g. 2026CS001."),
	("roll_number", "Roll Number", False, ""),
	("first_name", "First Name", True, ""),
	("last_name", "Last Name", False, ""),
	("gender", "Gender", True, "Male, Female or Other."),
	("date_of_birth", "Date of Birth", False, "DD-MM-YYYY."),
	("blood_group", "Blood Group", False, "A+, A-, B+, B-, AB+, AB-, O+ or O-."),
	("program", "Program", False, "Exact name of an Academic Program."),
	("year_of_study", "Year of Study", False, "1 to 5."),
	("admission_year", "Admission Year", False, "For example 2026."),
	(
		"academic_year",
		"Academic Year",
		False,
		"For example 2026-2027. Blank uses the import's academic year.",
	),
	("mobile_number", "Mobile Number", True, "10-digit mobile number; used for OTP login."),
	("email", "Email", False, ""),
	("address_line1", "Address Line 1", False, ""),
	("address_line2", "Address Line 2", False, ""),
	("city", "City", False, ""),
	("district", "District", False, ""),
	("state", "State", False, "Blank means Tamil Nadu."),
	("pincode", "PIN Code", False, "6 digits."),
	("emergency_contact_name", "Emergency Contact Name", False, ""),
	("emergency_contact_phone", "Emergency Contact Phone", False, "10-digit mobile number."),
	("mess_enrolled", "Mess Enrolled", False, "Yes or No. Blank means Yes."),
	("guardian_name", "Guardian Name", False, "Required when a guardian mobile is given."),
	("guardian_relation", "Guardian Relation", False, "Father, Mother, Guardian, Brother, Sister or Other."),
	("guardian_mobile", "Guardian Mobile", False, "Parent portal login number."),
	("guardian_email", "Guardian Email", False, ""),
	("guardian_occupation", "Guardian Occupation", False, ""),
	("guardian2_name", "Second Guardian Name", False, ""),
	("guardian2_relation", "Second Guardian Relation", False, ""),
	("guardian2_mobile", "Second Guardian Mobile", False, ""),
	("hostel", "Hostel", False, "Hostel code (e.g. BH1). Fill hostel and room to allot a bed."),
	("room_number", "Room Number", False, ""),
	("bed_label", "Bed", False, "Bed letter. Blank picks the first vacant bed in the room."),
	("allotment_from_date", "Allotment From Date", False, "DD-MM-YYYY. Blank means today."),
)
REQUIRED_COLUMNS = tuple(c[0] for c in COLUMNS if c[2])
COLUMN_ALIASES = {
	"reg_no": "register_number",
	"register_no": "register_number",
	"registration_number": "register_number",
	"roll_no": "roll_number",
	"mobile": "mobile_number",
	"mobile_no": "mobile_number",
	"phone": "mobile_number",
	"dob": "date_of_birth",
	"pin_code": "pincode",
	"pin": "pincode",
	"year": "year_of_study",
	"programme": "program",
	"bed_label": "bed_label",
	"bed": "bed_label",
	"room": "room_number",
	"parent_name": "guardian_name",
	"parent_mobile": "guardian_mobile",
}

GENDERS = {
	"male": "Male",
	"m": "Male",
	"boy": "Male",
	"female": "Female",
	"f": "Female",
	"girl": "Female",
	"other": "Other",
	"o": "Other",
	"transgender": "Other",
}
RELATIONS = ("Father", "Mother", "Guardian", "Brother", "Sister", "Other")
BLOOD_GROUPS = ("A+", "A-", "B+", "B-", "AB+", "AB-", "O+", "O-")
YES = {"yes", "y", "1", "true"}
NO = {"no", "n", "0", "false"}
DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d-%m-%y", "%d/%m/%y")

STUDENT_FIELDS = (
	"roll_number",
	"first_name",
	"last_name",
	"gender",
	"date_of_birth",
	"blood_group",
	"program",
	"year_of_study",
	"admission_year",
	"academic_year",
	"mobile_number",
	"email",
	"address_line1",
	"address_line2",
	"city",
	"district",
	"state",
	"pincode",
	"emergency_contact_name",
	"emergency_contact_phone",
)


class StudentBulkImport(Document):
	def validate(self):
		if self.has_value_changed("import_file") and not self.is_new():
			self.status = "Draft"
			self.validation_json = None
			self.total_rows = self.imported_rows = self.skipped_rows = self.failed_rows = 0
		if self.import_file and not self.import_file.lower().endswith((".csv", ".xlsx", ".xls")):
			frappe.throw(_("Attach an Excel (.xlsx) or CSV file."))

	# ------------------------------------------------------------------ actions

	@frappe.whitelist(methods=["POST"])
	def validate_file(self) -> dict:
		"""Check every row without writing anything and store the preview."""
		self.check_permission("write")
		if self.status == "Importing":
			frappe.throw(_("An import is already running."))

		rows, columns, ignored = self.read_rows()
		checker = RowChecker(self)
		results = [checker.check(row) for row in rows]

		errors = [r for r in results if r["errors"]]
		summary = {
			"rows": len(results),
			"create": sum(1 for r in results if not r["errors"] and r["action"] == "Create"),
			"update": sum(1 for r in results if not r["errors"] and r["action"] == "Update"),
			"skip": sum(1 for r in results if not r["errors"] and r["action"] == "Skip"),
			"errors": len(errors),
			"warnings": sum(1 for r in results if r["warnings"]),
			"allotments": sum(1 for r in results if not r["errors"] and r.get("bed")),
		}
		# The first rows give a feel for the file; every row with an error is listed so none is missed.
		shown = results[:PREVIEW_ROWS] + [r for r in results[PREVIEW_ROWS:] if r["errors"]]
		self.validation_json = json.dumps(
			{
				"columns": columns,
				"ignored_columns": ignored,
				"summary": summary,
				"rows": [_preview(r) for r in shown],
				"validated_on": str(now_datetime()),
			},
			default=str,
		)
		self.total_rows = len(results)
		self.status = "Validated"
		self.save()
		return summary

	@frappe.whitelist(methods=["POST"])
	def start_import(self) -> None:
		self.check_permission("write")
		frappe.has_permission("Hostel Student", "create", throw=True)
		if self.status == "Importing" and is_job_enqueued(self.job_id()):
			frappe.throw(_("This import is already running."))
		if self.status == "Draft":
			frappe.throw(_("Validate the file before importing."))

		self.status = "Importing"
		self.started_on = now_datetime()
		self.finished_on = None
		self.save()
		frappe.enqueue(
			"hostel360.hostel_people.doctype.student_bulk_import.student_bulk_import.run_import",
			queue="long",
			timeout=3600,
			job_id=self.job_id(),
			deduplicate=True,
			enqueue_after_commit=True,
			import_name=self.name,
		)

	def job_id(self) -> str:
		return f"hostel360-student-import::{self.name}"

	# ------------------------------------------------------------------ reading

	def get_file_doc(self):
		if not self.import_file:
			frappe.throw(_("Attach the file to import."))
		names = frappe.get_all(
			"File",
			filters={
				"file_url": self.import_file,
				"attached_to_doctype": self.doctype,
				"attached_to_name": self.name,
			},
			pluck="name",
			limit=1,
		) or frappe.get_all(
			"File",
			filters={"file_url": self.import_file, "owner": frappe.session.user},
			pluck="name",
			limit=1,
		)
		if not names:
			frappe.throw(_("The attached file was not found. Attach it to this import again."))
		return frappe.get_doc("File", names[0])

	def read_rows(self) -> tuple[list[dict], list[str], list[str]]:
		"""Rows as dicts keyed by template fieldname, plus recognised and ignored headers."""
		file_doc = self.get_file_doc()
		content = file_doc.get_content()
		extension = (file_doc.file_name or self.import_file).rsplit(".", 1)[-1].lower()
		try:
			if extension == "csv":
				from frappe.utils.csvutils import read_csv_content

				raw = read_csv_content(content)
			elif extension == "xlsx":
				from frappe.utils.xlsxutils import read_xlsx_file_from_attached_file

				raw = read_xlsx_file_from_attached_file(fcontent=content, read_only=True)
			elif extension == "xls":
				from frappe.utils.xlsxutils import read_xls_file_from_attached_file

				raw = read_xls_file_from_attached_file(content)
			else:
				frappe.throw(_("Attach an Excel (.xlsx) or CSV file."))
		except frappe.ValidationError:
			raise
		except Exception:
			frappe.log_error(
				title="Hostel360: unreadable student import file", message=frappe.get_traceback()
			)
			frappe.throw(_("The file could not be read. Save it again as .xlsx or .csv and attach it."))

		raw = [list(row) for row in (raw or []) if row and any(_cell(value) != "" for value in row)]
		if not raw:
			frappe.throw(_("The file is empty."))

		header = raw[0]
		mapping = {}
		ignored = []
		known = {c[0] for c in COLUMNS}
		labels = {_header_key(c[1]): c[0] for c in COLUMNS}
		for index, title in enumerate(header):
			key = _header_key(_cell(title))
			if not key:
				continue
			fieldname = key if key in known else labels.get(key) or COLUMN_ALIASES.get(key)
			if fieldname and fieldname not in mapping.values():
				mapping[index] = fieldname
			else:
				ignored.append(str(_cell(title)))

		missing = [
			label
			for fieldname, label, required, _help in COLUMNS
			if required and fieldname not in mapping.values()
		]
		if missing:
			frappe.throw(
				_(
					"These required columns are missing: {0}. Download the template to see the expected columns."
				).format(", ".join(_(m) for m in missing))
			)

		data = raw[1:]
		if len(data) > MAX_ROWS:
			frappe.throw(
				_("A file can hold at most {0} students. Split it into smaller files.").format(MAX_ROWS)
			)

		rows = []
		for offset, values in enumerate(data):
			row = {"row": offset + 2}
			for index, fieldname in mapping.items():
				row[fieldname] = _cell(values[index]) if index < len(values) else ""
			rows.append(row)
		return rows, list(mapping.values()), ignored

	# ------------------------------------------------------------------ importing

	def run_import(self) -> None:
		rows, _columns, _ignored = self.read_rows()
		checker = RowChecker(self)
		total = len(rows)
		log, failures = [], []
		imported = skipped = failed = 0

		for position, row in enumerate(rows, start=1):
			result = checker.check(row)
			label = f"Row {row['row']} · {result.get('register_number') or '-'}"
			if result["errors"]:
				failed += 1
				failures.append(
					{
						**_original(row),
						"errors": "; ".join(result["errors"]),
						"warnings": "; ".join(result["warnings"]),
					}
				)
				log.append(f"{label} · Failed: {'; '.join(result['errors'])}")
			elif result["action"] == "Skip":
				skipped += 1
				log.append(f"{label} · Skipped: {'; '.join(result['warnings']) or _('already registered')}")
			else:
				frappe.db.savepoint("student_import_row")
				try:
					outcome = self.import_row(result)
					frappe.db.release_savepoint("student_import_row")
					imported += 1
					notes = list(dict.fromkeys(result["warnings"] + outcome.get("notes", [])))
					log.append(f"{label} · {outcome['action']}" + (f" ({'; '.join(notes)})" if notes else ""))
					if result["warnings"]:
						failures.append(
							{**_original(row), "errors": "", "warnings": "; ".join(result["warnings"])}
						)
				except Exception as e:
					frappe.db.rollback(save_point="student_import_row")
					failed += 1
					message = _error_text(e)
					failures.append({**_original(row), "errors": message, "warnings": ""})
					log.append(f"{label} · Failed: {message}")
					frappe.clear_messages()
					checker.forget(result)

			if position % COMMIT_EVERY == 0:
				self.db_set(
					{"imported_rows": imported, "skipped_rows": skipped, "failed_rows": failed},
					update_modified=False,
				)
				frappe.db.commit()  # nosemgrep
			if position % 10 == 0 or position == total:
				frappe.publish_progress(
					position * 100 / total,
					title=_("Importing students"),
					doctype=self.doctype,
					docname=self.name,
					description=_("{0} of {1} rows").format(position, total),
				)

		if failed and not imported and not skipped:
			status = "Failed"
		elif failed:
			status = "Completed with Errors"
		else:
			status = "Completed"

		self.db_set(
			{
				"status": status,
				"total_rows": total,
				"imported_rows": imported,
				"skipped_rows": skipped,
				"failed_rows": failed,
				"finished_on": now_datetime(),
				"log": "\n".join(log),
				"errors_json": json.dumps(failures, default=str) if failures else None,
			}
		)
		self.notify_update()

	def import_row(self, result: dict) -> dict:
		values = result["values"]
		notes = []

		allotted = False
		if result.get("returning"):
			# A former resident of another hostel: the record becomes this hostel's to update
			# only once the student has a bed here, so the bed comes first.
			academic_year = values.get("academic_year") or frappe.db.get_value(
				"Hostel Student", result["student"], "academic_year"
			)
			self.allot_row(result["student"], result, academic_year, notes)
			allotted = True

		guardian_rows = []
		for guardian in result["guardians"]:
			guardian_rows.append((upsert_guardian(guardian, notes), guardian["relation"]))

		if result["student"]:
			student = frappe.get_doc("Hostel Student", result["student"])
			action = "Updated"
		else:
			student = frappe.new_doc("Hostel Student")
			student.register_number = result["register_number"]
			student.status = "Active"
			action = "Created"

		for fieldname in STUDENT_FIELDS:
			value = values.get(fieldname)
			if value not in (None, ""):
				student.set(fieldname, value)
		if values.get("mess_enrolled") is not None:
			student.mess_enrolled = values["mess_enrolled"]
		if student.is_new() and not student.state:
			student.state = "Tamil Nadu"

		linked = {row.guardian for row in student.guardians}
		has_primary = any(row.is_primary for row in student.guardians)
		for guardian, relation in guardian_rows:
			if guardian in linked:
				continue
			student.append(
				"guardians",
				{
					"guardian": guardian,
					"relation": relation,
					"is_primary": 0 if has_primary else 1,
					"can_approve_leave": 1,
				},
			)
			has_primary = True
			linked.add(guardian)

		student.save()

		if result.get("bed") and not allotted:
			self.allot_row(student.name, result, student.academic_year, notes)

		if self.create_portal_users:
			from hostel360.hostel_people.services import ensure_portal_user

			ensure_portal_user("Hostel Student", student.name)
			for guardian, _relation in guardian_rows:
				ensure_portal_user("Hostel Guardian", guardian)

		return {"action": action, "notes": notes}

	def allot_row(self, student: str, result: dict, academic_year: str | None, notes: list) -> None:
		allotment = frappe.get_doc(
			{
				"doctype": "Room Allotment",
				"student": student,
				"bed": result["bed"],
				"allotment_type": "New",
				"from_date": result.get("allotment_from_date") or nowdate(),
				"academic_year": academic_year,
				"remarks": _("Created by bulk import {0}").format(self.name),
			}
		)
		allotment.flags.skip_notifications = True
		allotment.insert()
		allotment.submit()
		notes.append(_("allotted bed {0}").format(result["bed"]))


# Guardian details an import may fill in when the registered guardian lacks them.
GUARDIAN_FILL_FIELDS = ("relation", "email", "occupation")


def upsert_guardian(guardian: dict, notes: list | None = None) -> str:
	"""Guardians are matched by mobile number; blank details are filled in, nothing is overwritten.

	A guardian who also has a child in a hostel the importing user does not look after is
	linked as registered: only staff who may edit every ward may change a guardian.
	"""
	name = frappe.db.get_value("Hostel Guardian", {"mobile_number": guardian["mobile_number"]}, "name")
	if name:
		doc = frappe.get_doc("Hostel Guardian", name)
		missing = guardian_fields_to_fill(doc, guardian)
		if not missing:
			return name
		if not frappe.has_permission("Hostel Guardian", "write", doc=doc):
			if notes is not None:
				notes.append(guardian_not_updated_note(guardian.get("label") or _("Guardian")))
			return name
		for fieldname in missing:
			doc.set(fieldname, guardian[fieldname])
		doc.save()
		return name

	doc = frappe.get_doc(
		{
			"doctype": "Hostel Guardian",
			"guardian_name": guardian["guardian_name"],
			"relation": guardian["relation"],
			"mobile_number": guardian["mobile_number"],
			"email": guardian.get("email"),
			"occupation": guardian.get("occupation"),
		}
	)
	doc.insert()
	return doc.name


def guardian_fields_to_fill(registered, guardian: dict) -> list[str]:
	return [field for field in GUARDIAN_FILL_FIELDS if guardian.get(field) and not registered.get(field)]


def guardian_not_updated_note(label: str) -> str:
	return _(
		"{0} linked as registered, without the missing details from the file: they also have a child in a hostel you do not look after"
	).format(label)


class RowChecker:
	"""Validates rows against the database and against earlier rows of the same file."""

	def __init__(self, import_doc: StudentBulkImport):
		self.doc = import_doc
		self.default_academic_year = import_doc.academic_year or frappe.db.get_single_value(
			"Hostel Settings", "current_academic_year"
		)
		self.programs = {name.lower(): name for name in frappe.get_all("Academic Program", pluck="name")}
		self.academic_years = set(frappe.get_all("Hostel Academic Year", pluck="name"))
		self.hostels = {}
		for row in frappe.get_list(
			"Hostel",
			filters={"is_active": 1},
			fields=["name", "hostel_name", "gender"],
			limit_page_length=0,
		):
			self.hostels[row.name.lower()] = row
			self.hostels[(row.hostel_name or "").lower()] = row
		self.can_allot = bool(
			frappe.has_permission("Room Allotment", "create")
			and frappe.has_permission("Room Allotment", "submit")
		)
		# Staff limited to some hostels or blocks may not change every registered student or
		# guardian; for everyone else the per-record checks below would only cost time.
		self.scoped = bool(
			get_user_permissions(frappe.session.user).get("Hostel") or block_restrictions(get_staff_scope())
		)
		self.seen_register_numbers = {}
		self.seen_mobiles = {}
		self.claimed_beds = set()
		self.room_cache = {}

	def forget(self, result: dict) -> None:
		"""A row that failed while importing releases the bed it had claimed."""
		if result.get("bed"):
			self.claimed_beds.discard(result["bed"])

	def check(self, row: dict) -> dict:
		errors, warnings = [], []
		values = {}
		result = {
			"row": row["row"],
			"register_number": None,
			"full_name": " ".join(
				filter(None, [str(row.get("first_name") or ""), str(row.get("last_name") or "")])
			).strip(),
			"student": None,
			"action": "Create",
			"values": values,
			"guardians": [],
			"bed": None,
			"errors": errors,
			"warnings": warnings,
		}

		def attempt(fn, *args):
			try:
				return fn(*args)
			except frappe.ValidationError as e:
				errors.append(_error_text(e))
				frappe.clear_messages()
				return None

		register_number = attempt(clean_register_number, str(row.get("register_number") or ""))
		result["register_number"] = register_number
		if register_number:
			if register_number in self.seen_register_numbers:
				errors.append(
					_("Register number repeats row {0}.").format(self.seen_register_numbers[register_number])
				)
			else:
				self.seen_register_numbers[register_number] = row["row"]

		existing = None
		if register_number:
			existing = frappe.db.get_value(
				"Hostel Student",
				register_number,
				["name", "gender", "current_hostel", "current_bed", "current_allotment", "academic_year"],
				as_dict=True,
			)
		if existing:
			result["student"] = existing.name
			result["action"] = "Update" if self.doc.update_existing else "Skip"
			if not self.doc.update_existing:
				warnings.append(_("Already registered; updating existing students is switched off."))
			elif self.scoped and not frappe.has_permission("Hostel Student", "write", doc=existing.name):
				if existing.current_hostel or existing.current_allotment:
					errors.append(
						_(
							"{0} lives in a hostel or block you do not look after. Ask that hostel's office or the Warden to update the record."
						).format(register_number)
					)
				else:
					# Last lived in another hostel. Allotting a bed here, in the same row, brings
					# the student to this hostel; the details are updated after that.
					result["returning"] = True

		for fieldname in (
			"roll_number",
			"first_name",
			"last_name",
			"address_line1",
			"address_line2",
			"city",
			"district",
			"state",
			"emergency_contact_name",
		):
			text = str(row.get(fieldname) or "").strip()
			if len(text) > 140:
				errors.append(_("{0} is too long.").format(_(_label(fieldname))))
			values[fieldname] = text

		if not values.get("first_name") and not existing:
			errors.append(_("First Name is required."))

		gender = GENDERS.get(str(row.get("gender") or "").strip().lower())
		if row.get("gender") and not gender:
			errors.append(_("Gender must be Male, Female or Other."))
		elif not gender and not existing:
			errors.append(_("Gender is required."))
		values["gender"] = gender

		mobile = None
		if row.get("mobile_number"):
			mobile = attempt(normalize_mobile, row.get("mobile_number"))
		elif not existing:
			errors.append(_("Mobile Number is required."))
		if mobile:
			owner = frappe.db.get_value("Hostel Student", {"mobile_number": mobile}, "name")
			if owner and owner != register_number:
				errors.append(_("Mobile number {0} belongs to student {1}.").format(mobile, owner))
			elif mobile in self.seen_mobiles:
				errors.append(_("Mobile number repeats row {0}.").format(self.seen_mobiles[mobile]))
			else:
				self.seen_mobiles[mobile] = row["row"]
		values["mobile_number"] = mobile

		email = str(row.get("email") or "").strip().lower()
		if email and not validate_email_address(email):
			errors.append(_("{0} is not a valid e-mail address.").format(email))
		values["email"] = email

		values["date_of_birth"] = self.parse_date(row.get("date_of_birth"), _("Date of Birth"), errors)
		if (
			values["date_of_birth"]
			and not 10 <= (getdate(nowdate()).year - values["date_of_birth"].year) <= 70
		):
			errors.append(
				_("Date of Birth {0} does not look right.").format(
					values["date_of_birth"].strftime("%d-%m-%Y")
				)
			)

		blood_group = re.sub(
			r"(VE|POSITIVE|NEGATIVE)$", "", str(row.get("blood_group") or "").upper().replace(" ", "")
		)
		blood_group = blood_group.replace("POS", "+").replace("NEG", "-")
		if blood_group and blood_group not in BLOOD_GROUPS:
			errors.append(_("Blood Group {0} is not recognised.").format(row.get("blood_group")))
		values["blood_group"] = blood_group or None

		program = str(row.get("program") or "").strip()
		if program:
			values["program"] = self.programs.get(program.lower())
			if not values["program"]:
				errors.append(_("Program {0} does not exist.").format(program))

		year = str(row.get("year_of_study") or "").strip()
		if year:
			if year not in ("1", "2", "3", "4", "5"):
				errors.append(_("Year of Study must be 1 to 5."))
			values["year_of_study"] = year

		admission_year = str(row.get("admission_year") or "").strip()
		if admission_year:
			if not admission_year.isdigit() or not 1950 <= int(admission_year) <= getdate(nowdate()).year + 1:
				errors.append(_("Admission Year {0} does not look right.").format(admission_year))
			else:
				values["admission_year"] = int(admission_year)

		academic_year = str(row.get("academic_year") or "").strip() or (
			None if existing and existing.academic_year else self.default_academic_year
		)
		if academic_year and academic_year not in self.academic_years:
			errors.append(_("Academic Year {0} does not exist.").format(academic_year))
		values["academic_year"] = academic_year

		pincode = str(row.get("pincode") or "").strip()
		if pincode:
			attempt(validate_pincode, pincode)
		values["pincode"] = pincode

		values["emergency_contact_phone"] = attempt(
			normalize_mobile, row.get("emergency_contact_phone") or None
		)

		mess = str(row.get("mess_enrolled") or "").strip().lower()
		if mess in YES:
			values["mess_enrolled"] = 1
		elif mess in NO:
			values["mess_enrolled"] = 0
		elif mess:
			errors.append(_("Mess Enrolled must be Yes or No."))
		else:
			values["mess_enrolled"] = None if existing else 1

		for prefix in ("guardian", "guardian2"):
			guardian = self.check_guardian(row, prefix, errors, warnings)
			if guardian:
				if guardian["mobile_number"] in {g["mobile_number"] for g in result["guardians"]}:
					errors.append(_("Both guardians have the same mobile number."))
				else:
					result["guardians"].append(guardian)

		if not errors and self.doc.allot_rooms and (row.get("hostel") or row.get("room_number")):
			self.check_allotment(row, result, existing, gender or (existing.gender if existing else None))

		if result.get("returning") and not errors:
			if not result.get("bed"):
				errors.append(
					_(
						"{0} last lived in another hostel, whose staff or the Warden keep the record. Give a hostel and room with a vacant bed in this row to bring the student to your hostel, or ask them to update it."
					).format(register_number)
				)
			elif gender and gender != existing.gender:
				# The bed is allotted for the registered gender before the details change.
				self.forget(result)
				result["bed"] = None
				errors.append(
					_(
						"Gender {0} differs from the registered {1}. Ask the Warden to correct it first."
					).format(_(gender), _(existing.gender or "-"))
				)
		return result

	def check_guardian(self, row: dict, prefix: str, errors: list, warnings: list) -> dict | None:
		name = " ".join(str(row.get(f"{prefix}_name") or "").split())
		raw_mobile = row.get(f"{prefix}_mobile")
		relation_text = str(row.get(f"{prefix}_relation") or "").strip()
		email = str(row.get(f"{prefix}_email") or "").strip().lower()
		occupation = str(row.get(f"{prefix}_occupation") or "").strip()
		label = _("Guardian") if prefix == "guardian" else _("Second guardian")

		if not (name or raw_mobile or relation_text or email or occupation):
			return None
		if not name or not raw_mobile:
			errors.append(_("{0}: name and mobile number are both required.").format(label))
			return None
		try:
			mobile = normalize_mobile(raw_mobile)
		except frappe.ValidationError as e:
			errors.append(f"{label}: {_error_text(e)}")
			frappe.clear_messages()
			return None

		relation = next((r for r in RELATIONS if r.lower() == relation_text.lower()), None)
		if relation_text and not relation:
			errors.append(_("{0}: relation {1} is not recognised.").format(label, relation_text))
		if email and not validate_email_address(email):
			errors.append(_("{0}: {1} is not a valid e-mail address.").format(label, email))
		guardian = {
			"label": label,
			"guardian_name": name,
			"mobile_number": mobile,
			"relation": relation or "Guardian",
			"email": email or None,
			"occupation": occupation or None,
		}
		self.check_registered_guardian(guardian, warnings)
		return guardian

	def check_registered_guardian(self, guardian: dict, warnings: list) -> None:
		"""Warn when the file would fill in a registered guardian the user may not change."""
		if not self.scoped:
			return
		registered = frappe.db.get_value(
			"Hostel Guardian",
			{"mobile_number": guardian["mobile_number"]},
			["name", *GUARDIAN_FILL_FIELDS],
			as_dict=True,
		)
		if not registered or not guardian_fields_to_fill(registered, guardian):
			return
		if not frappe.has_permission("Hostel Guardian", "write", doc=registered.name):
			warnings.append(guardian_not_updated_note(guardian["label"]))

	def check_allotment(self, row: dict, result: dict, existing, gender: str | None) -> None:
		warnings = result["warnings"]
		hostel = self.hostels.get(str(row.get("hostel") or "").strip().lower())
		room_number = re.sub(r"\s+", "", str(row.get("room_number") or "")).upper()
		bed_label = re.sub(r"\s+", "", str(row.get("bed_label") or "")).upper()

		def skip(message):
			warnings.append(_("Allotment skipped: {0}").format(message))

		if not self.can_allot:
			return skip(_("you are not permitted to allot rooms"))
		if not hostel:
			return skip(_("hostel {0} not found or not accessible").format(row.get("hostel") or "-"))
		if not room_number:
			return skip(_("room number is missing"))
		if hostel.gender in ("Male", "Female") and gender != hostel.gender:
			return skip(_("{0} is a {1} hostel").format(hostel.hostel_name, _(hostel.gender)))

		if existing and existing.current_allotment:
			current_room = frappe.db.get_value("Hostel Bed", existing.current_bed, "room")
			if current_room == f"{hostel.name}-{room_number}":
				return
			return skip(_("already allotted bed {0}; use Transfer").format(existing.current_bed))

		room_name = f"{hostel.name}-{room_number}"
		room = self.room_cache.get(room_name)
		if room is None:
			room = (
				frappe.db.get_value("Hostel Room", room_name, ["name", "status", "is_active"], as_dict=True)
				or {}
			)
			self.room_cache[room_name] = room
		if not room:
			return skip(_("room {0} does not exist in {1}").format(room_number, hostel.hostel_name))
		if not room.is_active or room.status in ("Under Maintenance", "Blocked"):
			return skip(_("room {0} is not open for allotment").format(room_number))

		if bed_label:
			bed = frappe.db.get_value(
				"Hostel Bed", f"{room_name}-{bed_label}", ["name", "status", "is_active"], as_dict=True
			)
			if not bed or not bed.is_active:
				return skip(_("bed {0} does not exist in room {1}").format(bed_label, room_number))
			if bed.status != "Vacant" or bed.name in self.claimed_beds:
				return skip(_("bed {0} of room {1} is not vacant").format(bed_label, room_number))
			chosen = bed.name
		else:
			vacant = frappe.get_all(
				"Hostel Bed",
				filters={"room": room_name, "status": "Vacant", "is_active": 1},
				fields=["name", "bed_label"],
			)
			vacant = sorted(
				(b for b in vacant if b.name not in self.claimed_beds),
				key=lambda b: (len(b.bed_label), b.bed_label),
			)
			if not vacant:
				return skip(_("room {0} has no vacant bed").format(room_number))
			chosen = vacant[0].name

		from_date = self.parse_date(
			row.get("allotment_from_date"), _("Allotment From Date"), result["errors"]
		)
		self.claimed_beds.add(chosen)
		result["bed"] = chosen
		result["allotment_from_date"] = from_date

	def parse_date(self, value, label: str, errors: list) -> date | None:
		if value in (None, ""):
			return None
		if isinstance(value, datetime):
			return value.date()
		if isinstance(value, date):
			return value
		text = str(value).strip()
		for fmt in DATE_FORMATS:
			try:
				return datetime.strptime(text, fmt).date()
			except ValueError:
				continue
		errors.append(_("{0} {1} is not a date (use DD-MM-YYYY).").format(label, text))
		return None


# ---------------------------------------------------------------------- background job


def run_import(import_name: str) -> None:
	doc = frappe.get_doc("Student Bulk Import", import_name)
	try:
		doc.run_import()
	except Exception as e:
		frappe.db.rollback()
		frappe.log_error(
			title="Hostel360: student import failed",
			message=frappe.get_traceback(),
			reference_doctype=doc.doctype,
			reference_name=doc.name,
		)
		doc.reload()
		doc.db_set(
			{
				"status": "Failed",
				"finished_on": now_datetime(),
				"log": (doc.log or "") + "\n" + _("Import stopped: {0}").format(_error_text(e)),
			}
		)
		doc.notify_update()


# ---------------------------------------------------------------------- downloads


@frappe.whitelist()
def download_template(file_format: str = "xlsx") -> None:
	frappe.has_permission("Student Bulk Import", "create", throw=True)
	headers = [label + (" *" if required else "") for _fieldname, label, required, _help in COLUMNS]

	if file_format == "csv":
		from frappe.utils.csvutils import to_csv

		frappe.response["filename"] = "hostel360-student-import-template.csv"
		frappe.response["filecontent"] = to_csv([headers]).encode("utf-8-sig")
		frappe.response["type"] = "binary"
		return

	frappe.response["filename"] = "hostel360-student-import-template.xlsx"
	frappe.response["filecontent"] = build_template_workbook(headers)
	frappe.response["type"] = "binary"


def build_template_workbook(headers: list[str]) -> bytes:
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.worksheet.datavalidation import DataValidation

	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "Students"
	sheet.append(headers)
	bold = Font(bold=True)
	required_fill = PatternFill("solid", fgColor="FFF4D6")
	for index, (_fieldname, label, required, _help) in enumerate(COLUMNS, start=1):
		cell = sheet.cell(row=1, column=index)
		cell.font = bold
		cell.alignment = Alignment(wrap_text=True, vertical="center")
		if required:
			cell.fill = required_fill
		sheet.column_dimensions[cell.column_letter].width = max(14, min(len(label) + 4, 30))
	sheet.freeze_panes = "A2"

	def add_list(fieldname, options):
		index = [c[0] for c in COLUMNS].index(fieldname) + 1
		letter = sheet.cell(row=1, column=index).column_letter
		validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
		validation.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
		sheet.add_data_validation(validation)

	add_list("gender", ["Male", "Female", "Other"])
	add_list("year_of_study", ["1", "2", "3", "4", "5"])
	add_list("blood_group", list(BLOOD_GROUPS))
	add_list("mess_enrolled", ["Yes", "No"])
	add_list("guardian_relation", list(RELATIONS))
	add_list("guardian2_relation", list(RELATIONS))

	# Text format keeps leading zeros and stops Excel turning mobile numbers into 9.88E+09.
	for fieldname in (
		"register_number",
		"roll_number",
		"mobile_number",
		"guardian_mobile",
		"guardian2_mobile",
		"emergency_contact_phone",
		"pincode",
		"room_number",
	):
		index = [c[0] for c in COLUMNS].index(fieldname) + 1
		letter = sheet.cell(row=1, column=index).column_letter
		for row in range(2, 502):
			sheet[f"{letter}{row}"].number_format = "@"

	guide = workbook.create_sheet("Instructions")
	guide.append(["Column", "Required", "What to enter"])
	for cell in guide[1]:
		cell.font = bold
	for _fieldname, label, required, help_text in COLUMNS:
		guide.append([label, "Yes" if required else "", help_text])
	guide.append([])
	guide.append(
		[
			"One student per row. Existing students are matched by register number and guardians by mobile number."
		]
	)
	guide.append(
		["Validate the file in Hostel360 before importing; every problem is listed against its row."]
	)
	guide.column_dimensions["A"].width = 28
	guide.column_dimensions["B"].width = 10
	guide.column_dimensions["C"].width = 90

	output = BytesIO()
	workbook.save(output)
	return output.getvalue()


@frappe.whitelist()
def download_error_report(name: str) -> None:
	doc = frappe.get_doc("Student Bulk Import", name)
	doc.check_permission("read")
	failures = json.loads(doc.errors_json or "[]")
	if not failures:
		frappe.throw(_("There are no failed rows in this import."))

	from frappe.utils.xlsxutils import make_xlsx

	columns = ["row"] + [c[0] for c in COLUMNS] + ["errors", "warnings"]
	labels = [_("Row")] + [c[1] for c in COLUMNS] + [_("Errors"), _("Warnings")]
	data = [labels] + [[failure.get(column, "") for column in columns] for failure in failures]
	frappe.response["filename"] = f"{doc.name}-rows-to-fix.xlsx"
	frappe.response["filecontent"] = make_xlsx(data, "Rows to fix").getvalue()
	frappe.response["type"] = "binary"


# ---------------------------------------------------------------------- helpers


def _cell(value):
	if value is None:
		return ""
	if isinstance(value, datetime):
		return value.date() if not (value.hour or value.minute or value.second) else value
	if isinstance(value, date):
		return value
	if isinstance(value, float) and value.is_integer():
		return str(int(value))
	return str(value).strip()


def _header_key(title) -> str:
	text = str(title or "").replace("*", "").strip().lower()
	return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def _label(fieldname: str) -> str:
	return next((c[1] for c in COLUMNS if c[0] == fieldname), fieldname)


def _error_text(error: Exception) -> str:
	text = str(error) or error.__class__.__name__
	return re.sub(r"<[^>]+>", "", text).strip()


def _original(row: dict) -> dict:
	return {
		key: (value.strftime("%d-%m-%Y") if isinstance(value, date) else value) for key, value in row.items()
	}


def _preview(result: dict) -> dict:
	return {
		"row": result["row"],
		"register_number": result.get("register_number"),
		"full_name": result.get("full_name"),
		"action": result["action"],
		"bed": result.get("bed"),
		"guardians": len(result.get("guardians") or []),
		"errors": result["errors"],
		"warnings": result["warnings"],
	}

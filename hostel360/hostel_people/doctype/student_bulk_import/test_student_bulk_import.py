# Copyright (c) 2026, HyFib Solutions Private Limited and contributors
# For license information, please see license.txt

import csv
import io
import json
from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase

from hostel360.hostel_people.doctype.student_bulk_import.student_bulk_import import (
	COLUMNS,
	build_template_workbook,
)
from hostel360.hostel_people.testing import (
	beds_of,
	make_hostel,
	make_room,
	make_staff,
	random_mobile,
	unique_code,
)

HEADER = [label + (" *" if required else "") for _fieldname, label, required, _help in COLUMNS]


def make_import(rows: list[dict], **values):
	buffer = io.StringIO()
	writer = csv.writer(buffer)
	writer.writerow(HEADER)
	for row in rows:
		writer.writerow([row.get(fieldname, "") for fieldname, _label, _required, _help in COLUMNS])

	file_doc = frappe.get_doc(
		{
			"doctype": "File",
			"file_name": f"students-{frappe.generate_hash(length=6)}.csv",
			"content": buffer.getvalue(),
			"is_private": 1,
		}
	).insert()
	return frappe.get_doc(
		{"doctype": "Student Bulk Import", "import_file": file_doc.file_url, **values}
	).insert()


def student_row(**values):
	register_number = values.pop("register_number", None) or unique_code("BI", 6)
	return {
		"register_number": register_number,
		"first_name": "Import",
		"last_name": register_number[-4:],
		"gender": "Male",
		"mobile_number": random_mobile(),
		"date_of_birth": "15-08-2007",
		"year_of_study": "1",
		**values,
	}


class TestStudentBulkImport(IntegrationTestCase):
	def setUp(self):
		self.hostel = make_hostel("Male")
		self.room = make_room(self.hostel.name, capacity=2, room_number="I1")

	def test_happy_path_creates_students_guardians_and_allotments(self):
		guardian_mobile = random_mobile()
		first = student_row(
			guardian_name="Senthil Kumar",
			guardian_relation="father",
			guardian_mobile=f"+91 {guardian_mobile}",
			hostel=self.hostel.name,
			room_number="i1",
			blood_group="b+ve",
			mess_enrolled="No",
		)
		second = student_row(gender="M", guardian_name="Senthil Kumar", guardian_mobile=guardian_mobile, hostel=self.hostel.name, room_number="I1")
		doc = make_import([first, second])

		summary = doc.validate_file()
		self.assertEqual(summary["rows"], 2)
		self.assertEqual(summary["create"], 2)
		self.assertEqual(summary["errors"], 0)
		self.assertEqual(summary["allotments"], 2)
		self.assertEqual(doc.status, "Validated")

		doc.run_import()
		doc.reload()
		self.assertEqual(doc.status, "Completed")
		self.assertEqual((doc.imported_rows, doc.failed_rows), (2, 0))

		student = frappe.get_doc("Hostel Student", first["register_number"])
		self.assertEqual(student.blood_group, "B+")
		self.assertEqual(student.mess_enrolled, 0)
		self.assertEqual(str(student.date_of_birth), "2007-08-15")
		self.assertEqual(len(student.guardians), 1)
		self.assertEqual(student.guardians[0].relation, "Father")
		self.assertEqual(student.current_bed, beds_of(self.room.name)[0])

		# Both students share the guardian, matched by mobile number.
		self.assertEqual(frappe.db.count("Hostel Guardian", {"mobile_number": guardian_mobile}), 1)
		other = frappe.get_doc("Hostel Student", second["register_number"])
		self.assertEqual(other.guardians[0].guardian, student.guardians[0].guardian)
		self.assertEqual(other.current_bed, beds_of(self.room.name)[1])
		self.assertEqual(frappe.db.get_value("Hostel Room", self.room.name, "status"), "Full")

		# Running the same file again updates in place.
		doc.validate_file()
		doc.run_import()
		doc.reload()
		self.assertEqual(doc.status, "Completed")
		self.assertEqual(frappe.db.count("Hostel Student", {"name": ["in", [first["register_number"], second["register_number"]]]}), 2)
		self.assertEqual(frappe.db.count("Room Allotment", {"student": first["register_number"], "docstatus": 1}), 1)

	def test_bad_rows_are_reported_and_good_rows_imported(self):
		good = student_row()
		duplicate = student_row(register_number=good["register_number"])
		rows = [
			good,
			student_row(mobile_number="12345"),
			student_row(gender=""),
			duplicate,
			student_row(program="No Such Program"),
			student_row(date_of_birth="31-31-2001"),
			student_row(guardian_name="Only Name"),
			student_row(email="not-an-email"),
		]
		doc = make_import(rows)
		summary = doc.validate_file()
		self.assertEqual(summary["rows"], 8)
		self.assertEqual(summary["errors"], 7)

		preview = json.loads(doc.validation_json)
		errors = {row["row"]: row["errors"] for row in preview["rows"] if row["errors"]}
		self.assertEqual(sorted(errors), [3, 4, 5, 6, 7, 8, 9])
		self.assertTrue(any("repeats row 2" in message for message in errors[5]))

		doc.run_import()
		doc.reload()
		self.assertEqual(doc.status, "Completed with Errors")
		self.assertEqual((doc.imported_rows, doc.failed_rows), (1, 7))
		self.assertTrue(frappe.db.exists("Hostel Student", good["register_number"]))
		self.assertEqual(len(json.loads(doc.errors_json)), 7)
		self.assertIn("Row 3", doc.log)

	def test_allotment_problems_do_not_block_the_student(self):
		female = student_row(gender="Female", hostel=self.hostel.name, room_number="I1")
		missing_room = student_row(hostel=self.hostel.name, room_number="NOPE")
		doc = make_import([female, missing_room])
		summary = doc.validate_file()
		self.assertEqual(summary["errors"], 0)
		self.assertEqual(summary["warnings"], 2)
		doc.run_import()
		doc.reload()
		self.assertEqual(doc.imported_rows, 2)
		self.assertFalse(frappe.db.get_value("Hostel Student", female["register_number"], "current_bed"))

	def test_missing_required_column_is_refused(self):
		file_doc = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": f"broken-{frappe.generate_hash(length=6)}.csv",
				"content": "Register Number,First Name\nX1,Test\n",
				"is_private": 1,
			}
		).insert()
		doc = frappe.get_doc({"doctype": "Student Bulk Import", "import_file": file_doc.file_url}).insert()
		self.assertRaises(frappe.ValidationError, doc.validate_file)

	def test_skip_existing_when_updates_are_off(self):
		row = student_row()
		first = make_import([row])
		first.validate_file()
		first.run_import()

		second = make_import([{**row, "first_name": "Renamed"}], update_existing=0)
		summary = second.validate_file()
		self.assertEqual(summary["skip"], 1)
		second.run_import()
		second.reload()
		self.assertEqual(second.skipped_rows, 1)
		self.assertEqual(frappe.db.get_value("Hostel Student", row["register_number"], "first_name"), "Import")

	def test_template_has_every_column(self):
		from openpyxl import load_workbook

		workbook = load_workbook(io.BytesIO(build_template_workbook(HEADER)))
		headers = [cell.value for cell in workbook["Students"][1]]
		self.assertEqual(headers, HEADER)
		self.assertIn("Instructions", workbook.sheetnames)

	def test_start_import_queues_a_long_job(self):
		doc = make_import([student_row()])
		self.assertRaises(frappe.ValidationError, doc.start_import)

		doc.validate_file()
		with patch("frappe.enqueue") as enqueue:
			doc.start_import()
		doc.reload()
		self.assertEqual(doc.status, "Importing")
		kwargs = enqueue.call_args.kwargs
		self.assertEqual(kwargs["queue"], "long")
		self.assertEqual(kwargs["import_name"], doc.name)
		self.assertTrue(kwargs["job_id"].endswith(doc.name))
		self.assertTrue(enqueue.call_args.args[0].endswith("student_bulk_import.run_import"))

	def test_scoped_staff_only_see_their_own_imports(self):
		# The attached spreadsheet carries students' and guardians' personal data.
		other_hostel = make_hostel("Male").name
		mine = make_staff("Office Staff", [self.hostel.name]).user
		theirs = make_staff("Deputy Warden", [other_hostel]).user
		warden = make_staff("Warden", []).user

		with self.set_user(mine):
			doc = make_import([student_row()])
			self.assertTrue(frappe.has_permission("Student Bulk Import", "read", doc=doc.name))
		with self.set_user(theirs):
			self.assertFalse(frappe.has_permission("Student Bulk Import", "read", doc=doc.name))
			self.assertNotIn(doc.name, frappe.get_list("Student Bulk Import", pluck="name"))
		with self.set_user(warden):
			self.assertTrue(frappe.has_permission("Student Bulk Import", "read", doc=doc.name))
